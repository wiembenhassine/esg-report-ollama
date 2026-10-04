"""
Export Word (.docx) des rapports, à partir du même Markdown que le HTML et le PDF.

La structure suit le rapport de référence (berkeley_gri_full.pdf) : titre et sous-titre,
« À propos de ce rapport », sections narratives, index de contenu GRI par norme, puis
annexes. Conversion volontairement simple et sans dépendance autre que python-docx :
titres, paragraphes (gras, italique, liens), listes, citations et tableaux. Les graphiques
SVG ne sont pas repris (leurs valeurs figurent déjà dans le tableau qui les précède).
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x2F, 0x6F, 0x5E)
TAG = re.compile(r"<[^>]+>")
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def plain(text: str) -> str:
    text = text.replace("<br>", "\n").replace("&nbsp;", " ").replace("\\|", "|")
    return TAG.sub("", text).strip()


def add_runs(paragraph, text: str) -> None:
    """Ajoute le texte au paragraphe en gérant **gras**, *italique*, `code` et [lien](url)."""
    for part in INLINE.split(plain(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            label, url = part[1:].split("](", 1)
            paragraph.add_run(f"{label} ({url[:-1]})")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def split_row(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip() for c in cells]


def add_table(doc, lines: list[str]) -> None:
    rows = [split_row(l) for l in lines if not re.match(r"^\|?\s*:?-{3,}", l.strip())]
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, row[j] if j < len(row) else "")
            for run in p.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def convert(md_text: str, path: Path, title: str) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    doc.core_properties.title = title

    lines = md_text.splitlines()
    i, para = 0, []

    def flush():
        if para:
            add_runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
        elif s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif s.startswith("<svg") or s.startswith('<p class="legend"'):
            flush()                                     # graphique : déjà donné en tableau
        elif s.startswith("<p"):
            flush()
            p = doc.add_paragraph()
            add_runs(p, s)
            for run in p.runs:
                run.italic = True
        elif m := re.match(r"^(#{1,4})\s+(.*)$", s):
            flush()
            level = len(m.group(1)) - 1                 # « # » = titre du document
            h = doc.add_heading(plain(m.group(2)), level=level)
            for run in h.runs:
                run.font.color.rgb = ACCENT
        elif s.startswith(("- ", "* ")):
            flush()
            add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif s.startswith(">"):
            flush()
            add_runs(doc.add_paragraph(style="Intense Quote"), s.lstrip("> "))
        else:
            para.append(s)
        i += 1
    flush()
    doc.save(path)
    return path
