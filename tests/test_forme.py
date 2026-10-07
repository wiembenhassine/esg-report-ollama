"""
Défauts de forme signalés par Wiem en lisant le rapport de Cork (6 octobre), corrigés au rendu :
1. ligne « :-: :-: » dans les tableaux Word ; 2. renvois GRI empilés ; 3. noms des universités ;
4. phrase « Elle est dénommée… » ; 5. limites répétées dans la section Social de Cork.
"""

import json
import re

from docx import Document

from esg import docx_export, guard, names, review
from esg.config import OUTPUTS


def corrected(key: str, section: str) -> str:
    d = json.loads((OUTPUTS / "cache" / key / f"{section}.json").read_text(encoding="utf-8"))
    text, _, _ = review.apply(key, section, d["text"])
    return names.normalize(key, guard.tidy(guard.drop_marker_echo(text), d["title"]))


def test_alignment_row_is_not_a_table_row_in_word(tmp_path):
    md = "# Titre\n\n| Rapporté | Non rapporté |\n|:-:|:-:|\n| 12 | 30 |\n\n| A | B |\n|---|--:|\n| x | y |\n"
    doc = Document(docx_export.convert(md, tmp_path / "t.docx", "Titre"))
    assert [len(t.rows) for t in doc.tables] == [2, 2]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert not [c for c in cells if set(c) <= set(":-") and c]


def test_stacked_gri_references_are_merged_on_one_line():
    text = ("Le plan a été élaboré après consultation.\n\n(GRI 2-22)\n\n(GRI 2-23)\n\n(GRI 2-24)\n\n"
            "### Limites et omissions\n\nCe rapport ne peut pas établir…\n\n(GRI 2-22), (GRI 2-23)")
    out = guard.merge_gri_refs(text)
    assert "(GRI 2-22, GRI 2-23, GRI 2-24)\n\n### Limites et omissions" in out
    assert out.endswith("(GRI 2-22, GRI 2-23)")
    single = "Texte.\n\n(GRI 2-1)\n\nAutre paragraphe."
    assert guard.merge_gri_refs(single) == single
    stacked = re.compile(r"\(GRI [^)]*\)\s*\n\s*\(GRI ")              # deux renvois GRI l'un sous l'autre
    for key in ("cork", "tudublin", "berkeley"):
        for section in ("organisation", "materialite", "gouvernance", "strategie", "parties_prenantes",
                        "environnement", "social"):
            assert not stacked.search(corrected(key, section)), f"{key}/{section}"


def test_university_names_full_then_short():
    cork = ("L'University College Cork (UCC) a obtenu un score. L'Université College Cork a aussi un plan, "
            "et la stratégie de l'University College Cork évolue.")
    assert names.normalize("cork", cork) == ("University College Cork (UCC) a obtenu un score. L'UCC a aussi "
                                             "un plan, et la stratégie de l'UCC évolue.")
    dublin = ("La Technological University Dublin (TU Dublin) a un plan. La TU Dublin estime que… "
              "Les salaires de la TU Dublin suivent les grilles. Le TU Dublin Strategic Plan le prévoit.")
    assert names.normalize("tudublin", dublin) == (
        "Technological University Dublin (TU Dublin) a un plan. TU Dublin estime que… Les salaires de TU Dublin "
        "suivent les grilles. Le TU Dublin Strategic Plan le prévoit.")
    berkeley = ("L'université de Californie, Berkeley, a déclaré son périmètre. Ce rapport ne peut pas établir le siège "
                "de l'université de Californie, Berkeley, car STARS ne le collecte pas. Le conseil des régents de "
                "l'université de Californie gouverne le système.")
    assert names.normalize("berkeley", berkeley) == (
        "University of California, Berkeley (UC Berkeley) a déclaré son périmètre. Ce rapport ne peut pas établir "
        "le siège de UC Berkeley, car STARS ne le collecte pas. Le conseil des régents de l'université de Californie "
        "gouverne le système.")
    for key in ("cork", "tudublin", "berkeley"):
        for section in ("organisation", "gouvernance", "strategie", "parties_prenantes", "social"):
            text = corrected(key, section)
            assert "L'University College Cork" not in text and "La TU Dublin" not in text
            assert "L'université de Californie à Berkeley" not in text


def test_cork_organisation_and_social_repetitions_are_gone():
    org = corrected("cork", "organisation")
    assert "Elle est dénommée" not in org
    social = corrected("cork", "social")
    assert "les données ne permettent pas d'établir" not in social
    assert social.count("Ce rapport ne peut pas établir") == 1
