"""Tests de la validation des données, des correspondances TCFD/ESRS, de l'export Word et de l'assistant."""

import csv

import pytest

from esg import chat, docx_export, frameworks, validate

FIELDS = ["institution", "stars_version", "pillar", "category", "category_code", "credit_code",
          "credit_name", "status", "score", "max"]
BERK = "University of California, Berkeley"


# ---------------------------------------------------------------- validation des données
def test_real_dataset_is_valid():
    checks = validate.validate()
    assert len(checks) == 191
    assert sum(c.kind == "erreur" for c in checks) == 0
    assert sum(c.kind == "noté" for c in checks) == 177
    assert sum(c.kind == "non noté" for c in checks) == 14


def test_invalid_rows_are_detected_and_excluded(tmp_path):
    rows = [
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-5", "Energy Use", "Complete", "5", "10"],   # ok
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-6", "GHG", "Complete", "17", "16"],        # > max
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-3", "Water", "Complete", "-1", "6"],       # < 0
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-4", "", "Complete", "2", "5"],             # nom vide
        ["Université X", "3.0", "Social", "Planning", "PA", "PA-8", "Gender", "Complete", "1", "2"],       # inconnue
        [BERK, "3.0", "Social", "Planning", "PA", "OP-9", "Mauvais code", "Complete", "1", "2"],           # incohérent
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-5", "Energy Use", "Complete", "5", "10"],   # doublon
        [BERK, "3.0", "Governance", "Planning", "PA", "PA-4", "Invest", "Not Applicable", "3", "4"],      # N.A. noté
        [BERK, "3.0", "Environmental", "Operations", "OP", "OP-7", "Dining", "Complete", "abc", "8"],     # non numérique
    ]
    path = tmp_path / "bad.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(FIELDS)
        w.writerows(rows)
    checks = validate.validate(path)
    assert [c.kind for c in checks] == ["noté"] + ["erreur"] * 8
    assert [r["credit_code"] for r in validate.valid_rows(path)] == ["OP-5"]
    assert "8 en erreur" in validate.report(checks)


# ---------------------------------------------------------------- correspondances
def test_every_stars_credit_is_mapped_by_hand():
    m = frameworks.load()
    for code in frameworks.credit_names():
        assert code in m["credits"], f"{code} absent de frameworks.yaml"
        for t in m["credits"][code]["tcfd"]:
            assert t in m["tcfd_labels"], f"{code}: code TCFD inconnu {t}"
        for e in m["credits"][code]["esrs"]:
            assert e in m["esrs_labels"], f"{code}: norme ESRS inconnue {e}"


def test_missing_correspondence_is_marked_not_forced():
    rows = {r["credit_code"]: r for r in frameworks.table()}
    assert rows["AC-1"]["tcfd"] == frameworks.NONE and rows["AC-1"]["esrs"] == frameworks.NONE
    assert rows["OP-6"]["tcfd"] != frameworks.NONE and "E1" in rows["OP-6"]["esrs"]
    assert rows["PA-13"]["tcfd"] == frameworks.NONE          # TCFD ne traite que du climat


# ---------------------------------------------------------------- export Word
def test_docx_export_keeps_structure(tmp_path):
    from docx import Document
    md = ("# Titre\n\n**À propos.** Texte avec *italique* et un [lien](https://x.org).\n\n## Section\n\n"
          "| A | B |\n|---|---|\n| <span class=\"badge ok\">Rapporté</span> | 1 / 2 |\n\n- point\n\n"
          '<svg viewBox="0 0 1 1"></svg>\n')
    path = docx_export.convert(md, tmp_path / "r.docx", "Titre")
    d = Document(path)
    assert d.paragraphs[0].style.name == "Title"
    assert any(p.style.name == "Heading 1" and p.text == "Section" for p in d.paragraphs)
    assert d.tables[0].cell(1, 0).text == "Rapporté"
    text = "\n".join(p.text for p in d.paragraphs)
    assert "**" not in text and "<span" not in text and "<svg" not in text


# ---------------------------------------------------------------- assistant
class NoIndex:
    def evidence(self, *a, **k):
        return []


def test_chat_question_without_data_is_answered_by_code(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("le modèle ne doit pas être appelé")
    monkeypatch.setattr(chat.llm, "chat", boom)
    r = chat.answer("Combien de prix Nobel a Berkeley ?")
    assert r["text"].startswith(chat.NOT_AVAILABLE) and r["source"] == "code"


def test_chat_numbers_come_from_code_and_model_numbers_are_removed(monkeypatch):
    monkeypatch.setattr(chat.rag, "get_index", lambda: NoIndex())
    reply = ("Cork obtient {{ cork_OP12_score }} sur {{ cork_OP12_max }} pour ses déchets, un niveau élevé. "
             "Cork recycle 80 % de ses déchets.")
    monkeypatch.setattr(chat.llm, "chat", lambda *a, **k: {"text": reply, "seconds": 0})
    r = chat.answer("Quel est le score de Cork sur la gestion des déchets ?")
    assert "3,99 sur 5" in r["text"]
    assert "80" not in r["text"] and len(r["removed"]) == 1


def test_chat_compares_universities_and_flags_physical_values(monkeypatch):
    monkeypatch.setattr(chat.rag, "get_index", lambda: NoIndex())
    monkeypatch.setattr(chat.llm, "chat", lambda *a, **k: {"text": "Berkeley obtient {{ berkeley_OP6_score }}, "
                        "TU Dublin {{ tudublin_OP6_score }}.", "seconds": 0})
    r = chat.answer("Combien de tonnes de CO2 émettent Berkeley et TU Dublin ?")
    assert r["keys"] == ["berkeley", "tudublin"]
    assert r["text"].startswith(chat.NOT_AVAILABLE)
    assert "8,01" in r["text"] and "10,79" in r["text"]
